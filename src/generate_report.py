#!/usr/bin/env python3

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from data_cleaning import clean_data
from kpi import build_summary, build_top_clients, build_source_performance

INPUT_FILE = "data/ventes_demo.csv"

OUTPUT_FILE = "output/reporting_auto.xlsx"



def format_excel(writer, sheet_names):
    """Applique une mise en forme simple aux onglets Excel."""

    workbook = writer.book

    header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")

    header_font = Font(bold=True)

    title_font = Font(bold=True, size=14)

    for sheet_name in sheet_names:

        ws = workbook[sheet_name]

        # Style des headers

        for cell in ws[1]:

            cell.fill = header_fill

            cell.font = header_font

            cell.alignment = Alignment(horizontal="center")

        # Ajustement largeur colonnes

        for column_cells in ws.columns:

            max_length = 0

            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(max_length, len(str(cell.value)))

            ws.column_dimensions[column_letter].width = min(max_length + 3, 35)

        # Figer la première ligne

        ws.freeze_panes = "A2"

        # Filtre automatique

        ws.auto_filter.ref = ws.dimensions

def main():

    # 1. Charger les données

    df_raw = pd.read_csv(INPUT_FILE)



    df_clean = clean_data(df_raw)

    # 3. Créer les tableaux

    summary = build_summary(df_clean)

    top_clients = build_top_clients(df_clean)

    source_performance = build_source_performance(df_clean)

    # 4. Exporter dans un seul fichier Excel avec 4 onglets

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:

        df_clean.to_excel(writer, sheet_name="Données nettoyées", index=False)

        summary.to_excel(writer, sheet_name="Résumé", index=False)

        top_clients.to_excel(writer, sheet_name="Top clients", index=False)

        source_performance.to_excel(writer, sheet_name="Performance sources", index=False)

        format_excel(

            writer,

            [

                "Données nettoyées",

                "Résumé",

                "Top clients",

                "Performance sources"

            ]

        )

    print(f"Reporting généré : {OUTPUT_FILE}")

if __name__ == "__main__":

    main()